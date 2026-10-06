from __future__ import annotations

import json
import math
from pathlib import Path
from urllib.request import urlopen

import matplotlib.pyplot as plt
from matplotlib.patches import Circle, Polygon
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
ASSETS.mkdir(exist_ok=True)

OUT = ROOT / "Militarisation_de_l_Arctique.pptx"
MAP_PATH = ASSETS / "carte_arctique.png"

W = 13.333
H = 7.5

NAVY = "071A2B"
NAVY_2 = "0B263B"
PANEL = "102F45"
PANEL_2 = "163C55"
ICE = "73D8FF"
CYAN = "2EBCE8"
BLUE = "3178C6"
WHITE = "F4F8FB"
MUTED = "AFC4D3"
MUTED_2 = "7895A8"
ORANGE = "FFB36B"
RED = "FF6B6B"
GREEN = "66D1A5"
YELLOW = "FFD166"
RUSSIA = "F06A6A"
LAND = "7892A4"
BLACK = "07131D"

FONT = "Inter"
FONT_FALLBACK = "Arial"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.replace("#", ""))


def add_rect(slide, x, y, w, h, fill, radius=True, line=None, line_width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_circle(slide, x, y, d, fill, line=None, line_width=1):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, x1, y1, x2, y2, color=MUTED_2, width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True

    p = frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    parts,
    x,
    y,
    w,
    h,
    size=18,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    for text, color, bold in parts:
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_paragraphs(
    slide,
    paragraphs,
    x,
    y,
    w,
    h,
    size=16,
    color=WHITE,
    bullet_color=ICE,
    gap=5,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = True
    for idx, item in enumerate(paragraphs):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r1 = p.add_run()
        r1.text = "• "
        r1.font.name = FONT
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = rgb(bullet_color)
        r2 = p.add_run()
        r2.text = item
        r2.font.name = FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = rgb(color)
    return box


def add_pill(slide, text, x, y, w, fill=PANEL_2, color=ICE, size=10):
    add_rect(slide, x, y, w, 0.31, fill, radius=True)
    add_text(
        slide,
        text.upper(),
        x,
        y + 0.015,
        w,
        0.27,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_header(slide, number, kicker, title, subtitle=None):
    add_pill(slide, kicker, 0.58, 0.38, max(1.2, len(kicker) * 0.075))
    add_text(slide, f"{number:02d}", 12.0, 0.37, 0.72, 0.3, 10, MUTED_2, True, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.58, 0.84, 11.9, 0.7, 29, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.48, 11.6, 0.42, 13, MUTED)


def add_footer(slide, source, number):
    add_line(slide, 0.58, 7.12, 12.75, 7.12, PANEL_2, 0.8)
    add_text(slide, source, 0.58, 7.17, 10.8, 0.2, 8, MUTED_2)
    add_text(slide, f"{number:02d}", 12.12, 7.16, 0.62, 0.2, 8, MUTED_2, True, align=PP_ALIGN.RIGHT)


def set_bg(slide, color=NAVY):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(color)


def add_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    frame.text = text.strip()
    for p in frame.paragraphs:
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(12)


def add_card_title(slide, number, title, x, y, w, accent=ICE):
    add_circle(slide, x, y, 0.35, accent)
    add_text(
        slide,
        str(number),
        x,
        y + 0.005,
        0.35,
        0.33,
        10,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(slide, title, x + 0.48, y - 0.01, w - 0.48, 0.4, 14, WHITE, True)


def polar_xy(lon, lat):
    max_r = 1.0
    r = (90.0 - lat) / 52.0 * max_r
    theta = math.radians(lon)
    return r * math.sin(theta), -r * math.cos(theta)


def country_class(name):
    nato = {
        "Canada",
        "United States of America",
        "Denmark",
        "Iceland",
        "Norway",
        "Sweden",
        "Finland",
    }
    if name == "Russia":
        return "russia"
    if name in nato:
        return "nato"
    return "other"


def fetch_geojson():
    url = (
        "https://raw.githubusercontent.com/nvkelso/"
        "natural-earth-vector/master/geojson/ne_110m_admin_0_countries.geojson"
    )
    with urlopen(url, timeout=30) as response:
        return json.load(response)


def draw_geo_ring(ax, ring, face, edge):
    chunks = []
    chunk = []
    last_lon = None
    for lon, lat in ring:
        if lat < 34:
            if len(chunk) >= 3:
                chunks.append(chunk)
            chunk = []
            last_lon = None
            continue
        if last_lon is not None and abs(lon - last_lon) > 180:
            if len(chunk) >= 3:
                chunks.append(chunk)
            chunk = []
        chunk.append(polar_xy(lon, lat))
        last_lon = lon
    if len(chunk) >= 3:
        chunks.append(chunk)
    for points in chunks:
        ax.add_patch(
            Polygon(points, closed=True, facecolor=face, edgecolor=edge, lw=0.55, zorder=4)
        )


def make_map():
    data = fetch_geojson()
    fig, ax = plt.subplots(figsize=(8, 8), dpi=220)
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")

    ax.add_patch(Circle((0, 0), 1.015, facecolor="#0B2E43", edgecolor="#2B6682", lw=1.2))
    for lat, lw in [(80, 0.5), (70, 0.55), (60, 0.6), (50, 0.65)]:
        r = (90 - lat) / 52
        ax.add_patch(Circle((0, 0), r, fill=False, edgecolor="#2A5268", lw=lw, alpha=0.85))
    for lon in range(0, 360, 30):
        x, y = polar_xy(lon, 38)
        ax.plot([0, x], [0, y], color="#24485C", lw=0.45, zorder=1)

    fills = {"nato": "#2EBCE8", "russia": "#F06A6A", "other": "#7892A4"}
    for feature in data["features"]:
        name = feature["properties"].get("NAME", "")
        geom = feature["geometry"]
        if geom is None:
            continue
        cls = country_class(name)
        face = fills[cls]
        edge = "#D3E8F2" if cls != "other" else "#AABEC9"
        coords = geom["coordinates"]
        polygons = coords if geom["type"] == "MultiPolygon" else [coords]
        for polygon in polygons:
            if not polygon:
                continue
            draw_geo_ring(ax, polygon[0], face, edge)

    # Northern Sea Route, schematic.
    nsr = [
        (32, 70),
        (55, 73),
        (85, 76),
        (115, 76),
        (145, 73),
        (170, 67),
    ]
    pts = [polar_xy(lon, lat) for lon, lat in nsr]
    ax.plot(
        [p[0] for p in pts],
        [p[1] for p in pts],
        color="#FFD166",
        lw=2.1,
        zorder=8,
    )

    locations = {
        "KOLA": (34, 69, "#FFFFFF"),
        "GROENLAND": (-42, 72, "#FFFFFF"),
        "BÉRING": (178, 66, "#FFFFFF"),
    }
    for label, (lon, lat, color) in locations.items():
        x, y = polar_xy(lon, lat)
        ax.scatter([x], [y], s=15, c=color, zorder=10)
        ax.text(
            x + 0.035,
            y + 0.025,
            label,
            fontsize=6.8,
            color=color,
            fontweight="bold",
            family="DejaVu Sans",
            zorder=10,
        )

    ax.text(
        0.17,
        -0.42,
        "Route maritime du Nord",
        fontsize=6.5,
        color="#FFD166",
        family="DejaVu Sans",
        zorder=10,
        rotation=-16,
    )

    ax.set_xlim(-1.08, 1.08)
    ax.set_ylim(-1.08, 1.08)
    ax.set_aspect("equal")
    ax.axis("off")
    plt.tight_layout(pad=0)
    fig.savefig(MAP_PATH, transparent=True, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)


def add_chevron(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    return shape


def add_source_link(slide, label, title, url, x, y, w):
    add_text(slide, label, x, y, 0.36, 0.28, 10, ICE, True)
    title_box = add_text(slide, title, x + 0.43, y - 0.02, w - 0.43, 0.25, 10, WHITE, True)
    url_box = add_text(slide, url, x + 0.43, y + 0.24, w - 0.43, 0.31, 8, MUTED_2)
    for box in (title_box, url_box):
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.hyperlink.address = url


def build_presentation():
    make_map()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1 — Cover
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_rect(slide, 7.35, -0.35, 6.5, 8.2, NAVY_2, radius=False)
    slide.shapes.add_picture(str(MAP_PATH), Inches(7.58), Inches(0.18), width=Inches(5.45))
    add_pill(slide, "GÉOPOLITIQUE · 2026", 0.62, 0.56, 2.02)
    add_text(slide, "LA MILITARISATION", 0.62, 1.35, 6.65, 0.72, 35, WHITE, True)
    add_text(slide, "DE L’ARCTIQUE", 0.62, 2.06, 6.6, 0.8, 38, ICE, True)
    add_text(
        slide,
        "La fonte des glaces transforme-t-elle le Grand Nord\nen nouveau front stratégique ?",
        0.65,
        3.15,
        5.95,
        1.02,
        18,
        MUTED,
    )
    add_line(slide, 0.65, 4.55, 2.25, 4.55, ICE, 3)
    add_text(slide, "EXPOSÉ", 0.65, 4.82, 1.25, 0.34, 11, MUTED_2, True)
    add_text(slide, "Défense · climat · rivalités de puissance", 0.65, 5.22, 5.8, 0.4, 15, WHITE, True)
    add_text(slide, "Carte schématique — frontières non destinées à l’arbitrage", 8.1, 6.83, 4.55, 0.22, 7, MUTED_2, align=PP_ALIGN.RIGHT)
    add_notes(
        slide,
        """
        Bonjour. Mon exposé porte sur la militarisation de l’Arctique.
        La question n’est pas seulement de savoir si les armées y sont plus présentes,
        mais pourquoi elles le sont et quels risques cette présence crée.
        Mon idée principale est la suivante : le climat rend l’espace plus accessible,
        mais c’est la rivalité entre puissances qui transforme cette accessibilité en enjeu militaire.
        """,
    )

    # 2 — Drivers
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        2,
        "LE DÉCLENCHEUR",
        "Un espace qui change de fonction",
        "Trois dynamiques se renforcent — sans produire mécaniquement une guerre.",
    )
    cards = [
        (
            "01",
            "ACCESSIBILITÉ",
            "4,60 M km²",
            "minimum de banquise en septembre 2025",
            "−12,1 % par décennie\nsur 1979–2025",
            ICE,
        ),
        (
            "02",
            "VALEUR STRATÉGIQUE",
            "Routes",
            "saisonnières, ports, énergie et minerais",
            "Mais des coûts extrêmes\net une logistique fragile",
            ORANGE,
        ),
        (
            "03",
            "SÉCURITÉ",
            "3 fonctions",
            "dissuasion nucléaire · alerte · projection",
            "L’Arctique relie\nAtlantique et Pacifique",
            RED,
        ),
    ]
    for i, (num, title, stat, desc, foot, accent) in enumerate(cards):
        x = 0.62 + i * 4.12
        add_rect(slide, x, 2.1, 3.78, 3.83, PANEL, radius=True)
        add_pill(slide, num, x + 0.26, 2.36, 0.62, fill=accent, color=NAVY, size=9)
        add_text(slide, title, x + 0.26, 2.83, 3.25, 0.32, 11, accent, True)
        add_text(slide, stat, x + 0.26, 3.35, 3.25, 0.58, 27, WHITE, True)
        add_text(slide, desc, x + 0.26, 4.02, 3.18, 0.68, 13, MUTED)
        add_line(slide, x + 0.26, 4.92, x + 3.42, 4.92, PANEL_2, 1)
        add_text(slide, foot, x + 0.26, 5.15, 3.18, 0.53, 12, WHITE, True)
    add_rect(slide, 0.62, 6.18, 12.08, 0.63, NAVY_2, radius=True, line=PANEL_2)
    add_rich_text(
        slide,
        [
            ("IDÉE-CLÉ  ", ICE, True),
            ("Le climat ouvre l’espace ; ", WHITE, False),
            ("la rivalité le militarise.", ORANGE, True),
        ],
        0.92,
        6.35,
        11.52,
        0.3,
        15,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Source : NSIDC, minimum annuel de la banquise, 2025.", 2)
    add_notes(
        slide,
        """
        Premier facteur : l’accessibilité. En 2025, le minimum de banquise atteint 4,60 millions
        de kilomètres carrés. La tendance de long terme reste une baisse de 12,1 % par décennie,
        même si la baisse n’est pas linéaire chaque année.
        Deuxième facteur : l’espace intéresse pour les routes, les ressources et les infrastructures.
        Mais il reste très coûteux et dangereux.
        Enfin, l’Arctique a trois fonctions militaires : protéger la dissuasion nucléaire,
        surveiller les approches et projeter des forces entre Atlantique et Pacifique.
        """,
    )

    # 3 — Strategic map
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        3,
        "LA CARTE DU RAPPORT DE FORCE",
        "8 États arctiques — désormais 7 alliés face à la Russie",
        "L’élargissement de l’OTAN transforme la géographie stratégique du Nord.",
    )
    add_rect(slide, 0.55, 2.0, 7.45, 4.88, NAVY_2, radius=True)
    slide.shapes.add_picture(str(MAP_PATH), Inches(0.82), Inches(2.1), width=Inches(4.85))
    add_pill(slide, "OTAN", 5.45, 2.42, 0.86, fill=CYAN, color=NAVY)
    add_pill(slide, "RUSSIE", 6.4, 2.42, 1.1, fill=RUSSIA, color=NAVY)
    add_text(slide, "Route maritime du Nord", 5.5, 2.94, 2.05, 0.4, 10, YELLOW, True)
    add_text(slide, "schématique", 5.5, 3.25, 1.7, 0.27, 8, MUTED_2)
    side = [
        ("7 / 8", "États arctiques\nsont membres de l’OTAN", ICE),
        ("1 340 km", "frontière terrestre\nFinlande–Russie", ORANGE),
        ("0 côte", "pour la Chine,\nacteur non riverain", GREEN),
    ]
    for i, (stat, desc, accent) in enumerate(side):
        y = 2.08 + i * 1.47
        add_rect(slide, 8.35, y, 4.35, 1.22, PANEL, radius=True)
        add_text(slide, stat, 8.68, y + 0.19, 1.4, 0.45, 23, accent, True)
        add_text(slide, desc, 10.05, y + 0.2, 2.32, 0.62, 12, WHITE, True)
    add_rect(slide, 8.35, 6.5, 4.35, 0.38, PANEL_2, radius=True)
    add_text(
        slide,
        "Un arc allié quasi continu : Alaska → Finlande",
        8.48,
        6.57,
        4.05,
        0.2,
        9,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(
        slide,
        "Sources : OTAN (adhésions Finlande 2023, Suède 2024) ; carte Natural Earth, adaptée.",
        3,
    )
    add_notes(
        slide,
        """
        L’Arctique compte huit États : Canada, États-Unis, Danemark via le Groenland,
        Islande, Norvège, Suède, Finlande et Russie.
        Depuis l’entrée de la Finlande en 2023 et de la Suède en 2024,
        sept de ces huit États sont membres de l’OTAN.
        L’Alliance dispose donc d’un arc géographique presque continu de l’Alaska à la Finlande.
        La Russie reste toutefois le plus grand État arctique et possède la plus longue façade.
        La Chine n’a aucune côte arctique : son influence dépend donc de partenariats et d’accès.
        """,
    )

    # 4 — Russia
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        4,
        "LA POSTURE RUSSE",
        "Le « bastion » du Nord : protéger pour pouvoir frapper",
        "La péninsule de Kola concentre dissuasion nucléaire et déni d’accès.",
    )
    add_rect(slide, 0.62, 2.02, 6.12, 4.72, NAVY_2, radius=True)
    cx, cy = 3.68, 4.36
    add_circle(slide, cx - 1.77, cy - 1.77, 3.54, NAVY_2, line=PANEL_2, line_width=2)
    add_circle(slide, cx - 1.25, cy - 1.25, 2.5, PANEL, line=RUSSIA, line_width=2)
    add_circle(slide, cx - 0.71, cy - 0.71, 1.42, RUSSIA)
    add_text(slide, "KOLA", cx - 0.52, cy - 0.32, 1.04, 0.28, 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "SNLE", cx - 0.45, cy + 0.02, 0.9, 0.2, 9, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "DÉFENSE AÉRIENNE · MISSILES CÔTIERS", 1.42, 2.57, 4.52, 0.27, 9, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "SOUS-MARINS · AVIATION · CAPTEURS", 1.42, 5.86, 4.52, 0.27, 9, MUTED, True, align=PP_ALIGN.CENTER)
    add_chevron(slide, 5.78, 4.02, 0.62, 0.72, RUSSIA)
    add_text(slide, "VERS\nL’ATLANTIQUE", 5.34, 4.9, 1.52, 0.54, 8, MUTED_2, True, align=PP_ALIGN.CENTER)
    points = [
        ("DISSUASION", "Protéger les sous-marins nucléaires stratégiques — capacité de seconde frappe."),
        ("DÉNI D’ACCÈS", "S-300/S-400, missiles côtiers, radars et aviation en couches successives."),
        ("ROUTE DU NORD", "Bases et capteurs soutiennent contrôle, secours et souveraineté revendiquée."),
        ("DOUBLE USAGE", "Une posture défensive pour Moscou peut menacer les renforts alliés."),
    ]
    for i, (title, desc) in enumerate(points):
        y = 2.08 + i * 1.12
        add_rect(slide, 7.12, y, 5.58, 0.92, PANEL, radius=True)
        add_text(slide, f"{i + 1:02d}", 7.39, y + 0.21, 0.46, 0.24, 10, RED, True)
        add_text(slide, title, 7.96, y + 0.13, 1.55, 0.23, 10, WHITE, True)
        add_text(slide, desc, 9.48, y + 0.11, 2.91, 0.63, 10, MUTED)
    add_rect(slide, 7.12, 6.57, 5.58, 0.2, RUSSIA, radius=True)
    add_footer(slide, "Sources : CSIS, The Ice Curtain ; Russian Arctic Threat.", 4)
    add_notes(
        slide,
        """
        La priorité russe n’est pas seulement de conquérir des territoires.
        Elle consiste d’abord à protéger la flotte du Nord et les sous-marins nucléaires
        basés autour de la péninsule de Kola. C’est le concept de bastion :
        plusieurs couches de défense aérienne, missiles côtiers, aviation et sous-marins
        créent un espace protégé. Le paradoxe est classique :
        ce que Moscou présente comme défensif peut menacer les routes de renfort de l’OTAN
        dans l’Atlantique Nord. C’est le cœur du dilemme de sécurité.
        """,
    )

    # 5 — NATO timeline
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        5,
        "LA RÉPONSE ALLIÉE",
        "L’OTAN se « nordifie » en quatre étapes",
        "L’intégration politique devient une architecture opérationnelle.",
    )
    add_line(slide, 1.07, 3.82, 12.25, 3.82, PANEL_2, 4)
    timeline = [
        ("2022", "RUPTURE", "Guerre en Ukraine\nCoopération politique gelée", RED),
        ("2023", "FINLANDE", "Adhésion à l’OTAN\nFrontière +1 340 km", ICE),
        ("2024", "SUÈDE", "Adhésion à l’OTAN\nNordiques réunis", CYAN),
        ("2025–26", "INTÉGRATION", "JFC Norfolk · FLF Finland\nArctic Sentry", ORANGE),
    ]
    xs = [1.07, 4.02, 6.97, 9.92]
    for i, ((year, title, desc, accent), x) in enumerate(zip(timeline, xs)):
        add_circle(slide, x, 3.55, 0.54, accent)
        add_text(slide, str(i + 1), x, 3.65, 0.54, 0.22, 9, NAVY, True, align=PP_ALIGN.CENTER)
        card_y = 2.02 if i % 2 == 0 else 4.28
        add_rect(slide, x - 0.25, card_y, 2.6, 1.28, PANEL, radius=True)
        add_text(slide, year, x, card_y + 0.16, 1.04, 0.33, 16, accent, True)
        add_text(slide, title, x + 1.04, card_y + 0.18, 1.0, 0.26, 9, WHITE, True, align=PP_ALIGN.RIGHT)
        add_text(slide, desc, x, card_y + 0.61, 2.06, 0.48, 10, MUTED)
        if i % 2 == 0:
            add_line(slide, x + 0.27, card_y + 1.28, x + 0.27, 3.55, accent, 1.4)
        else:
            add_line(slide, x + 0.27, 4.09, x + 0.27, card_y, accent, 1.4)
    add_rect(slide, 0.65, 6.28, 12.0, 0.55, NAVY_2, radius=True, line=PANEL_2)
    add_rich_text(
        slide,
        [
            ("ARCTIC SENTRY  ", ORANGE, True),
            ("activité multidomaine dirigée par JFC Norfolk ; exercices et vigilance réunis.", WHITE, False),
        ],
        0.9,
        6.44,
        11.5,
        0.26,
        12,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Sources : OTAN, 11 février et 8 juin 2026 ; JFC Norfolk, 5 décembre 2025.", 5)
    add_notes(
        slide,
        """
        La réponse de l’OTAN se lit en quatre étapes.
        En 2022, l’invasion de l’Ukraine provoque une rupture politique majeure.
        La Finlande adhère en 2023, puis la Suède en 2024.
        À partir de 2025, cette nouvelle carte est transformée en commandement intégré :
        les pays nordiques rejoignent la zone de responsabilité de JFC Norfolk.
        En 2026, Arctic Sentry rassemble exercices, surveillance et activités multidomaines.
        Une force avancée multinationale pour la Finlande est également mise en place.
        """,
    )

    # 6 — Friction points
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        6,
        "LES ZONES SENSIBLES",
        "Trois espaces de friction — trois logiques différentes",
        "Le risque vient moins d’une bataille unique que d’incidents reliés entre eux.",
    )
    friction = [
        (
            "01",
            "KOLA · BARENTS · GIUK",
            "Dissuasion et Atlantique",
            ["Sous-marins nucléaires", "Lutte anti-sous-marine", "Renforts et câbles"],
            "RISQUE : mauvaise interprétation d’une opération",
            RED,
        ),
        (
            "02",
            "GROENLAND · AMÉRIQUE",
            "Alerte et espace",
            ["Radar et défense antimissile", "Souveraineté groenlandaise", "Accès aux infrastructures"],
            "RISQUE : pression politique et dual-use",
            ICE,
        ),
        (
            "03",
            "BÉRING · ROUTE DU NORD",
            "Contrôle et coopération russo-chinoise",
            ["Régulation russe du passage", "Patrouilles et exercices", "Ports et énergie"],
            "RISQUE : zone grise, collision, brouillage",
            ORANGE,
        ),
    ]
    for i, (num, title, sub, bullets, risk, accent) in enumerate(friction):
        x = 0.62 + i * 4.12
        add_rect(slide, x, 2.12, 3.78, 4.56, PANEL, radius=True)
        add_pill(slide, num, x + 0.25, 2.38, 0.62, fill=accent, color=NAVY)
        add_text(slide, title, x + 0.25, 2.89, 3.2, 0.34, 11, WHITE, True)
        add_text(slide, sub, x + 0.25, 3.31, 3.2, 0.46, 11, accent, True)
        for j, bullet in enumerate(bullets):
            by = 4.03 + j * 0.48
            add_circle(slide, x + 0.27, by + 0.06, 0.12, accent)
            add_text(slide, bullet, x + 0.53, by, 2.9, 0.3, 11, MUTED)
        add_rect(slide, x + 0.24, 5.71, 3.3, 0.68, NAVY_2, radius=True)
        add_text(slide, risk, x + 0.42, 5.86, 2.94, 0.36, 9, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Sources : DoD, Arctic Strategy 2024 ; CSIS, 2025.", 6)
    add_notes(
        slide,
        """
        Trois espaces concentrent les risques.
        Autour de Kola, l’enjeu est la dissuasion nucléaire et l’accès à l’Atlantique par le passage GIUK,
        entre Groenland, Islande et Royaume-Uni.
        Autour du Groenland et de l’Amérique du Nord, les priorités sont l’alerte avancée,
        l’espace et les infrastructures duales, à la fois civiles et militaires.
        Enfin, vers le détroit de Béring et la route maritime du Nord,
        la Russie veut contrôler le passage tandis que la coopération avec la Chine augmente.
        La présence chinoise reste cependant bien plus limitée que celle de la Russie ou des États-Unis.
        """,
    )

    # 7 — China
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        7,
        "LE TROISIÈME ACTEUR",
        "La Chine : puissance polaire sans côte arctique",
        "Une influence réelle, mais une empreinte militaire encore limitée.",
    )
    add_rect(slide, 0.62, 2.05, 5.62, 4.72, PANEL, radius=True)
    add_pill(slide, "AMBITION", 0.93, 2.37, 1.12, fill=ORANGE, color=NAVY)
    add_text(slide, "2018", 0.93, 2.9, 1.48, 0.6, 31, ORANGE, True)
    add_text(slide, "Pékin se décrit comme\n« État proche de l’Arctique »", 2.28, 2.88, 3.42, 0.65, 15, WHITE, True)
    ambitions = [
        ("SCIENCE", "stations, recherche, données"),
        ("ÉCONOMIE", "GNL, minerais, infrastructures"),
        ("ROUTES", "« Route polaire de la soie »"),
        ("PARTENARIAT", "énergie et activités avec la Russie"),
    ]
    for i, (title, desc) in enumerate(ambitions):
        y = 3.9 + i * 0.58
        add_text(slide, title, 0.93, y, 1.28, 0.24, 10, ORANGE, True)
        add_text(slide, desc, 2.28, y, 3.48, 0.28, 11, MUTED)

    add_rect(slide, 6.55, 2.05, 6.15, 4.72, NAVY_2, radius=True, line=PANEL_2)
    add_pill(slide, "LIMITES", 6.86, 2.37, 1.0, fill=ICE, color=NAVY)
    limits = [
        ("Aucune côte arctique", "l’accès dépend d’États riverains"),
        ("Peu de bases permanentes", "empreinte militaire non comparable"),
        ("Milieu très contraignant", "glace, distance, secours, coûts"),
        ("Relation asymétrique", "Moscou coopère sans céder le contrôle"),
    ]
    for i, (title, desc) in enumerate(limits):
        y = 2.98 + i * 0.76
        add_circle(slide, 6.88, y + 0.05, 0.22, ICE)
        add_text(slide, title, 7.25, y, 2.08, 0.27, 11, WHITE, True)
        add_text(slide, desc, 9.35, y, 2.92, 0.34, 10, MUTED)
    add_rect(slide, 6.87, 6.13, 5.5, 0.38, PANEL, radius=True)
    add_text(
        slide,
        "ACTEUR D’INFLUENCE ≠ PUISSANCE MILITAIRE DOMINANTE",
        7.0,
        6.22,
        5.25,
        0.2,
        9,
        ICE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Sources : Livre blanc chinois sur l’Arctique, 2018 ; DoD, 2024 ; CSIS, 2025.", 7)
    add_notes(
        slide,
        """
        La Chine constitue un troisième acteur, mais il faut éviter de l’exagérer.
        En 2018, elle se qualifie elle-même d’État proche de l’Arctique.
        Ses intérêts sont scientifiques, économiques et maritimes.
        La coopération avec la Russie s’étend dans l’énergie et certaines activités de sécurité.
        Mais la Chine n’a aucune côte arctique, peu de bases permanentes et dépend des États riverains.
        Elle est donc aujourd’hui surtout un acteur d’influence et d’opportunité,
        pas une puissance militaire arctique au niveau de la Russie ou de l’OTAN.
        """,
    )

    # 8 — Escalation loop
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        8,
        "LE DILEMME DE SÉCURITÉ",
        "Militarisation ne signifie pas guerre inévitable",
        "Le danger le plus crédible : l’escalade involontaire.",
    )
    add_rect(slide, 0.62, 2.05, 7.25, 4.72, PANEL, radius=True)
    add_text(slide, "LA BOUCLE D’ESCALADE", 0.93, 2.37, 3.1, 0.32, 11, RED, True)
    stages = [
        ("1", "ACCÈS", "climat · routes"),
        ("2", "DÉPLOIEMENTS", "bases · exercices"),
        ("3", "PERCEPTION", "menace supposée"),
        ("4", "RÉACTION", "contre-mesures"),
    ]
    coords = [(1.05, 3.15), (3.72, 3.15), (3.72, 5.12), (1.05, 5.12)]
    for (num, title, sub), (x, y) in zip(stages, coords):
        add_rect(slide, x, y, 2.28, 1.04, NAVY_2, radius=True, line=PANEL_2)
        add_circle(slide, x + 0.18, y + 0.2, 0.34, RED)
        add_text(slide, num, x + 0.18, y + 0.27, 0.34, 0.18, 9, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.67, y + 0.17, 1.36, 0.25, 10, WHITE, True)
        add_text(slide, sub, x + 0.67, y + 0.51, 1.42, 0.24, 9, MUTED)
    add_chevron(slide, 3.34, 3.42, 0.3, 0.46, RED)
    add_chevron(slide, 4.68, 4.53, 0.3, 0.46, RED)
    add_chevron(slide, 3.34, 5.39, 0.3, 0.46, RED)
    add_chevron(slide, 1.9, 4.53, 0.3, 0.46, RED)
    add_circle(slide, 6.13, 3.55, 1.16, RED)
    add_text(slide, "INCIDENT", 6.22, 3.92, 0.98, 0.25, 10, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "erreur · collision\nbrouillage", 5.98, 4.85, 1.47, 0.6, 10, MUTED, True, align=PP_ALIGN.CENTER)

    add_rect(slide, 8.16, 2.05, 4.54, 4.72, NAVY_2, radius=True, line=PANEL_2)
    add_text(slide, "LES FREINS", 8.49, 2.37, 2.2, 0.32, 11, GREEN, True)
    brakes = [
        ("01", "Dissuasion", "coût d’un affrontement direct"),
        ("02", "Environnement", "logistique lente et vulnérable"),
        ("03", "Droit maritime", "frontières largement délimitées"),
        ("04", "Coopération ciblée", "science, secours, contacts"),
    ]
    for i, (num, title, desc) in enumerate(brakes):
        y = 2.96 + i * 0.75
        add_text(slide, num, 8.49, y, 0.4, 0.23, 9, GREEN, True)
        add_text(slide, title, 8.98, y, 1.45, 0.25, 11, WHITE, True)
        add_text(slide, desc, 10.45, y, 1.88, 0.35, 9, MUTED)
    add_rect(slide, 8.48, 6.1, 3.88, 0.38, PANEL, radius=True)
    add_text(slide, "Conseil de l’Arctique : sécurité exclue du mandat", 8.6, 6.2, 3.62, 0.2, 8, MUTED, True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Sources : Conseil de l’Arctique, 2025–2026 ; analyse de synthèse.", 8)
    add_notes(
        slide,
        """
        La militarisation crée un dilemme de sécurité.
        Un État déploie des moyens qu’il juge défensifs ; l’autre les interprète comme offensifs
        et répond à son tour. Dans un milieu difficile, une collision, une erreur de navigation,
        un brouillage ou une alerte mal comprise peuvent accélérer cette boucle.
        Mais la guerre n’est pas inévitable : la dissuasion, les coûts logistiques,
        le droit maritime et la coopération technique imposent des limites.
        Le Conseil de l’Arctique continue certains travaux scientifiques,
        mais son mandat exclut précisément les questions militaires.
        """,
    )

    # 9 — 2030 scenarios
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        9,
        "PROJECTION",
        "Quatre scénarios à l’horizon 2030",
        "Probabilité et impact ne se confondent pas.",
    )
    add_rect(slide, 0.62, 2.05, 8.15, 4.7, NAVY_2, radius=True)
    x0, y0, x1, y1 = 1.35, 6.05, 8.1, 2.62
    add_line(slide, x0, y0, x1, y0, MUTED_2, 1.3)
    add_line(slide, x0, y0, x0, y1, MUTED_2, 1.3)
    add_text(slide, "PROBABILITÉ →", 6.62, 6.2, 1.48, 0.25, 9, MUTED_2, True, align=PP_ALIGN.RIGHT)
    add_text(slide, "IMPACT ↑", 0.81, 2.47, 0.44, 1.04, 9, MUTED_2, True, align=PP_ALIGN.CENTER)
    add_line(slide, 4.72, 2.63, 4.72, 6.05, PANEL_2, 1)
    add_line(slide, 1.35, 4.34, 8.1, 4.34, PANEL_2, 1)
    scenarios = [
        ("CONFLIT OUVERT", 2.08, 3.03, 1.55, 0.68, RED, "faible prob. · impact extrême"),
        ("ZONE GRISE", 5.85, 3.5, 1.62, 0.68, ORANGE, "incident · sabotage · brouillage"),
        ("COOPÉRATION SÉLECTIVE", 2.15, 5.02, 2.02, 0.68, GREEN, "science · secours · pêche"),
        ("DISSUASION GÉRÉE", 5.47, 4.82, 1.95, 0.68, ICE, "scénario central"),
    ]
    for title, x, y, w, h, color, detail in scenarios:
        add_rect(slide, x, y, w, h, color, radius=True)
        add_text(slide, title, x + 0.08, y + 0.09, w - 0.16, 0.21, 8, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.08, y + 0.34, w - 0.16, 0.22, 7, NAVY, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 9.1, 2.05, 3.6, 4.7, PANEL, radius=True)
    add_pill(slide, "LECTURE", 9.42, 2.37, 0.9, fill=ICE, color=NAVY)
    add_text(slide, "Le plus probable", 9.42, 2.96, 2.95, 0.29, 11, ICE, True)
    add_text(slide, "Une dissuasion tendue mais contenue.", 9.42, 3.31, 2.76, 0.58, 14, WHITE, True)
    add_line(slide, 9.42, 4.1, 12.36, 4.1, PANEL_2, 1)
    add_text(slide, "Le plus dangereux", 9.42, 4.36, 2.95, 0.29, 11, RED, True)
    add_text(slide, "Un incident de zone grise qui touche une capacité nucléaire.", 9.42, 4.71, 2.76, 0.78, 14, WHITE, True)
    add_text(slide, "Jugement analytique — pas une prévision certaine.", 9.42, 6.17, 2.82, 0.28, 8, MUTED_2)
    add_footer(slide, "Scénarios construits à partir des tendances observées en 2024–2026.", 9)
    add_notes(
        slide,
        """
        Pour 2030, quatre scénarios peuvent être distingués.
        Le scénario central est une dissuasion gérée : plus d’exercices et de surveillance,
        mais pas de conflit direct.
        La coopération sélective peut survivre dans la science, le secours ou la pêche.
        Les actions de zone grise — sabotage, brouillage, pression sur des infrastructures —
        sont assez plausibles et plus dangereuses.
        Le conflit ouvert reste le moins probable, mais son impact serait extrême,
        surtout si une capacité nucléaire était concernée.
        """,
    )

    # 10 — Recommendations
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        10,
        "ÉVITER L’ENGRENAGE",
        "Quatre principes pour une sécurité soutenable",
        "Dissuader sans rendre chaque mouvement illisible.",
    )
    recommendations = [
        ("01", "DISSUASION CRÉDIBLE", "Présence proportionnée, défense aérienne, capacités de secours.", ICE),
        ("02", "TRANSPARENCE", "Notification d’exercices, canaux de crise, règles d’interception.", GREEN),
        ("03", "RÉSILIENCE", "Ports, satellites, câbles et navigation protégés contre les zones grises.", ORANGE),
        ("04", "SÉCURITÉ HUMAINE", "Climat et populations autochtones intégrés aux décisions.", YELLOW),
    ]
    positions = [(0.62, 2.09), (6.75, 2.09), (0.62, 4.4), (6.75, 4.4)]
    for (num, title, desc, accent), (x, y) in zip(recommendations, positions):
        add_rect(slide, x, y, 5.95, 1.92, PANEL, radius=True)
        add_circle(slide, x + 0.3, y + 0.3, 0.58, accent)
        add_text(slide, num, x + 0.3, y + 0.48, 0.58, 0.22, 10, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 1.08, y + 0.31, 4.37, 0.3, 13, WHITE, True)
        add_text(slide, desc, x + 1.08, y + 0.79, 4.27, 0.65, 12, MUTED)
        add_rect(slide, x + 1.08, y + 1.55, 3.15, 0.08, accent, radius=True)
    add_footer(slide, "Synthèse : OTAN, Conseil de l’Arctique, DoD et littérature stratégique.", 10)
    add_notes(
        slide,
        """
        Une stratégie efficace doit combiner quatre principes.
        Premièrement, une dissuasion crédible, avec des capacités adaptées au milieu.
        Deuxièmement, de la transparence : notification des exercices, canaux de crise
        et règles claires lors des interceptions.
        Troisièmement, la résilience des ports, câbles, satellites et systèmes de navigation.
        Enfin, la sécurité ne peut pas être seulement militaire :
        le changement climatique et les populations autochtones doivent être intégrés aux décisions.
        """,
    )

    # 11 — Conclusion
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        11,
        "CONCLUSION",
        "Oui, l’Arctique se militarise — mais de façon asymétrique",
        "La région est d’abord un espace de dissuasion, de surveillance et de gestion du risque.",
    )
    takeaways = [
        ("01", "RUSSIE", "Avantage géographique et forte densité nucléaire autour de Kola.", RED),
        ("02", "OTAN", "Avantage de coalition et continuité territoriale depuis 2024.", ICE),
        ("03", "CHINE", "Influence croissante, mais empreinte militaire encore secondaire.", ORANGE),
    ]
    for i, (num, actor, message, accent) in enumerate(takeaways):
        y = 2.12 + i * 1.18
        add_rect(slide, 0.65, y, 7.42, 0.94, PANEL, radius=True)
        add_circle(slide, 0.93, y + 0.22, 0.49, accent)
        add_text(slide, num, 0.93, y + 0.36, 0.49, 0.19, 9, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, actor, 1.66, y + 0.21, 1.22, 0.26, 11, accent, True)
        add_text(slide, message, 2.89, y + 0.2, 4.72, 0.47, 12, WHITE, True)
    add_rect(slide, 8.43, 2.12, 4.27, 3.3, NAVY_2, radius=True, line=PANEL_2)
    add_text(slide, "RÉPONSE À LA PROBLÉMATIQUE", 8.76, 2.48, 3.62, 0.3, 10, ICE, True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "La fonte des glaces\ncrée l’opportunité.\n\nLa rivalité stratégique\ncrée la militarisation.",
        8.82,
        3.02,
        3.5,
        1.7,
        20,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, 0.65, 6.13, 12.05, 0.63, PANEL_2, radius=True)
    add_text(
        slide,
        "Le défi n’est pas de supprimer toute présence militaire — mais d’empêcher qu’elle devienne incontrôlable.",
        0.97,
        6.32,
        11.42,
        0.26,
        13,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Conclusion de l’exposé.", 11)
    add_notes(
        slide,
        """
        Pour conclure, oui, l’Arctique se militarise, mais de manière asymétrique.
        La Russie conserve l’avantage de la géographie et une forte densité nucléaire.
        L’OTAN dispose désormais de l’avantage de la coalition et d’une continuité territoriale.
        La Chine influence l’équilibre à long terme, sans dominer militairement aujourd’hui.
        La fonte des glaces crée donc l’opportunité, tandis que la rivalité stratégique crée la militarisation.
        Le véritable défi est d’empêcher que la présence militaire devienne incontrôlable.
        """,
    )

    # 12 — Sources
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(
        slide,
        12,
        "SOURCES",
        "Références essentielles",
        "Documents officiels et analyses consultés — liens cliquables.",
    )
    left_sources = [
        (
            "01",
            "OTAN — Arctic Sentry, 11 fév. 2026",
            "https://www.nato.int/en/news-and-events/articles/news/2026/02/11/nato-secretary-general-outlines-new-activity-arctic-sentry-ahead-of-defence-ministers-meeting",
        ),
        (
            "02",
            "OTAN — sécurité du Grand Nord, 8 juin 2026",
            "https://www.nato.int/en/news-and-events/articles/news/2026/06/08/nato-enhances-security-in-the-arctic-and-high-north",
        ),
        (
            "03",
            "JFC Norfolk — intégration des Alliés nordiques, 2025",
            "https://jfcnorfolk.nato.int/activity/joint-force-command-norfolk-welcomes-nordic-allies-to-its-area-of-responsibility",
        ),
        (
            "04",
            "NSIDC — minimum de banquise 2025",
            "https://nsidc.org/sea-ice-today/analyses/2025-arctic-sea-ice-minimum-squeezes-ten-lowest-minimums",
        ),
    ]
    right_sources = [
        (
            "05",
            "Conseil de l’Arctique — mandat et organisation",
            "https://arctic-council.org/about/",
        ),
        (
            "06",
            "Conseil de l’Arctique — mise à jour 2026",
            "https://arctic-council.org/news/2026-arctic-council-update/",
        ),
        (
            "07",
            "CSIS — The Ice Curtain: Kola Peninsula",
            "https://www.csis.org/analysis/ice-curtain-modernization-kola-peninsula",
        ),
        (
            "08",
            "DoD — 2024 Arctic Strategy",
            "https://media.defense.gov/2024/Jul/22/2003507411/-1/-1/0/DOD-ARCTIC-STRATEGY-2024.PDF",
        ),
    ]
    for i, item in enumerate(left_sources):
        add_source_link(slide, *item, 0.68, 2.13 + i * 1.03, 5.86)
    for i, item in enumerate(right_sources):
        add_source_link(slide, *item, 6.78, 2.13 + i * 1.03, 5.84)
    add_rect(slide, 0.68, 6.34, 11.9, 0.4, NAVY_2, radius=True)
    add_text(
        slide,
        "Carte : Natural Earth, adaptée · Données vérifiées au 6 octobre 2026",
        0.9,
        6.44,
        11.45,
        0.19,
        8,
        MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Les notes du présentateur contiennent un script oral pour chaque diapositive.", 12)
    add_notes(
        slide,
        """
        Voici les principales sources utilisées.
        J’ai privilégié les documents officiels de l’OTAN, du Conseil de l’Arctique,
        du Département américain de la Défense et du NSIDC,
        complétés par les analyses du CSIS.
        Les liens sont cliquables dans la présentation.
        Merci pour votre attention.
        """,
    )

    # Core properties
    props = prs.core_properties
    props.title = "La militarisation de l’Arctique"
    props.subject = "Exposé géopolitique — situation en 2026"
    props.author = "Cursor"
    props.keywords = "Arctique, militarisation, OTAN, Russie, Chine, géopolitique"
    props.comments = "Présentation en français avec notes du présentateur et sources cliquables."

    prs.save(OUT)
    print(f"Created: {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_presentation()
