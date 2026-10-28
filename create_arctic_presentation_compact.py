from pathlib import Path

from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from create_arctic_presentation import (
    CYAN,
    GREEN,
    ICE,
    MAP_PATH,
    MUTED,
    MUTED_2,
    NAVY,
    NAVY_2,
    ORANGE,
    PANEL,
    PANEL_2,
    RED,
    RUSSIA,
    WHITE,
    YELLOW,
    add_circle,
    add_footer,
    add_line,
    add_notes,
    add_pill,
    add_rect,
    add_rich_text,
    add_text,
    make_map,
    rgb,
)


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Militarisation_Arctique_Version_Courte.pptx"

HERO = ASSETS / "arctic_icebreaker_hero.jpg"
QUESTIONS = ASSETS / "arctic_questions_background.jpg"
HERO_DARK = ASSETS / "arctic_icebreaker_hero_dark.jpg"
QUESTIONS_DARK = ASSETS / "arctic_questions_background_dark.jpg"

W = 13.333
H = 7.5


def prepare_image(source: Path, target: Path, brightness: float, navy_overlay: int):
    image = Image.open(source).convert("RGB")
    image = ImageEnhance.Brightness(image).enhance(brightness)
    overlay = Image.new("RGB", image.size, f"#{NAVY}")
    image = Image.blend(image, overlay, navy_overlay / 100)
    image = image.filter(ImageFilter.GaussianBlur(radius=0.18))
    image.save(target, quality=94, optimize=True)


def prepare_assets():
    if not HERO.exists() or not QUESTIONS.exists():
        raise FileNotFoundError(
            "Les images sources doivent être placées dans assets/ avant génération."
        )
    prepare_image(HERO, HERO_DARK, 0.68, 20)
    prepare_image(QUESTIONS, QUESTIONS_DARK, 0.58, 28)
    make_map()


def set_bg(slide, color=NAVY):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_full_image(slide, path):
    slide.shapes.add_picture(str(path), 0, 0, width=Inches(W), height=Inches(H))


def add_slide_number(slide, number):
    add_text(
        slide,
        f"{number:02d}",
        12.23,
        0.38,
        0.5,
        0.22,
        9,
        MUTED,
        True,
        align=PP_ALIGN.RIGHT,
    )


def add_kicker(slide, number, text):
    add_pill(slide, f"{number:02d} · {text}", 0.62, 0.42, max(1.62, len(text) * 0.08 + 0.74))
    add_slide_number(slide, number)


def add_photo_credit(slide, text):
    add_text(slide, text, 9.55, 6.92, 3.18, 0.18, 7, MUTED_2, align=PP_ALIGN.RIGHT)


def add_icon_route(slide, x, y, color):
    add_circle(slide, x, y, 0.14, color)
    add_line(slide, x + 0.14, y + 0.07, x + 0.88, y + 0.07, color, 2.2)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x + 0.78), Inches(y - 0.05), Inches(0.26), Inches(0.25)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()


def add_icon_radar(slide, x, y, color):
    add_circle(slide, x + 0.31, y + 0.31, 0.16, color)
    for d in (0.5, 0.82, 1.14):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ARC,
            Inches(x + 0.39 - d / 2),
            Inches(y + 0.39 - d / 2),
            Inches(d),
            Inches(d),
        )
        shape.fill.background()
        shape.line.color.rgb = rgb(color)
        shape.line.width = Inches(0.018)


def build():
    prepare_assets()

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1 — Hero cover
    slide = prs.slides.add_slide(blank)
    add_full_image(slide, HERO_DARK)
    add_rect(slide, 0, 0, 6.55, 7.5, NAVY, radius=False)
    add_pill(slide, "EXPOSÉ · 2026", 0.68, 0.62, 1.55)
    add_text(slide, "LA MILITARISATION", 0.68, 1.72, 5.9, 0.62, 32, WHITE, True)
    add_text(slide, "DE L’ARCTIQUE", 0.68, 2.35, 5.8, 0.72, 38, ICE, True)
    add_line(slide, 0.7, 3.42, 2.2, 3.42, ICE, 3)
    add_text(
        slide,
        "Climat  ·  puissance  ·  sécurité",
        0.7,
        3.78,
        4.8,
        0.38,
        15,
        WHITE,
        True,
    )
    add_text(
        slide,
        "Une nouvelle frontière stratégique ?",
        0.7,
        4.42,
        4.95,
        0.42,
        17,
        MUTED,
    )
    add_text(slide, "FORMAT COURT", 0.7, 6.65, 2.0, 0.22, 9, MUTED_2, True)
    add_photo_credit(slide, "Illustration éditoriale générée · sans valeur documentaire")
    add_notes(
        slide,
        """
        Bonjour. Aujourd’hui, je vais présenter la militarisation de l’Arctique.
        La question centrale est simple : la fonte des glaces transforme-t-elle
        cette région en nouvelle frontière stratégique ?
        Nous allons voir pourquoi l’Arctique attire les puissances,
        qui y renforce sa présence, puis où se situent les principaux risques.
        """,
    )

    # 2 — Why now?
    slide = prs.slides.add_slide(blank)
    add_full_image(slide, HERO_DARK)
    add_rect(slide, 0, 0, 7.42, 7.5, NAVY, radius=False)
    add_kicker(slide, 2, "POURQUOI MAINTENANT ?")
    add_text(slide, "L’espace s’ouvre.", 0.65, 1.2, 6.2, 0.62, 31, WHITE, True)
    add_text(slide, "La compétition suit.", 0.65, 1.82, 6.2, 0.62, 31, ICE, True)
    add_text(
        slide,
        "Le climat rend la région plus accessible —\npas plus facile à contrôler.",
        0.67,
        2.74,
        5.85,
        0.72,
        15,
        MUTED,
    )

    cards = [
        ("−12,1 %", "BANQUISE", "par décennie", ICE),
        ("ROUTES", "ACCÈS", "saisonnier", ORANGE),
        ("ALERTE", "DÉFENSE", "nucléaire", RED),
    ]
    for i, (big, title, small, color) in enumerate(cards):
        x = 0.66 + i * 2.16
        add_rect(slide, x, 4.05, 1.92, 1.78, PANEL, radius=True)
        add_circle(slide, x + 0.22, 4.3, 0.2, color)
        add_text(slide, big, x + 0.22, 4.69, 1.52, 0.33, 16, WHITE, True)
        add_text(slide, title, x + 0.22, 5.13, 1.48, 0.22, 9, color, True)
        add_text(slide, small, x + 0.22, 5.42, 1.48, 0.2, 9, MUTED)

    add_rect(slide, 8.17, 4.78, 4.1, 1.06, NAVY_2, radius=True)
    add_rich_text(
        slide,
        [
            ("CLIMAT ", ICE, True),
            ("→ ", MUTED, True),
            ("ACCÈS ", ORANGE, True),
            ("→ ", MUTED, True),
            ("PUISSANCE", RED, True),
        ],
        8.46,
        5.15,
        3.5,
        0.27,
        13,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Source : NSIDC, minimum de banquise 2025.", 2)
    add_photo_credit(slide, "Illustration générée")
    add_notes(
        slide,
        """
        Pourquoi maintenant ? D’abord à cause du climat.
        Depuis 1979, le minimum de banquise de septembre diminue de 12,1 % par décennie.
        Cela rend certaines routes et certaines zones plus accessibles pendant une partie de l’année.
        Mais l’Arctique reste très difficile : froid, distance, manque de ports et secours limités.
        Son importance militaire vient surtout de l’alerte avancée,
        des sous-marins nucléaires et de la liaison entre Atlantique et Pacifique.
        """,
    )

    # 3 — Actors
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_kicker(slide, 3, "QUI MILITARISE ?")
    add_text(slide, "Un face-à-face asymétrique", 0.62, 1.16, 7.0, 0.56, 29, WHITE, True)
    add_rich_text(
        slide,
        [("7 / 8 ", ICE, True), ("États arctiques sont membres de l’OTAN", MUTED, False)],
        0.65,
        1.78,
        5.6,
        0.34,
        14,
    )
    add_rect(slide, 0.55, 2.27, 7.13, 4.57, NAVY_2, radius=True)
    slide.shapes.add_picture(str(MAP_PATH), Inches(0.68), Inches(2.3), width=Inches(4.58))
    add_pill(slide, "OTAN", 5.18, 2.58, 0.82, fill=CYAN, color=NAVY)
    add_pill(slide, "RUSSIE", 6.09, 2.58, 1.02, fill=RUSSIA, color=NAVY)
    add_text(
        slide,
        "Carte schématique",
        5.18,
        3.08,
        1.95,
        0.22,
        8,
        MUTED_2,
    )

    actors = [
        ("RUSSIE", "Kola · SNLE · déni d’accès", RED, "BASTION"),
        ("OTAN", "continuité · surveillance · renforts", ICE, "ALLIANCE"),
        ("CHINE", "science · routes · influence", ORANGE, "ACCÈS"),
    ]
    for i, (name, line, color, tag) in enumerate(actors):
        y = 2.28 + i * 1.42
        add_rect(slide, 8.03, y, 4.68, 1.18, PANEL, radius=True)
        add_circle(slide, 8.35, y + 0.28, 0.55, color)
        add_text(
            slide,
            name[0],
            8.35,
            y + 0.43,
            0.55,
            0.2,
            10,
            NAVY,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, name, 9.13, y + 0.22, 1.2, 0.28, 13, color, True)
        add_pill(slide, tag, 11.25, y + 0.19, 1.04, fill=NAVY_2, color=MUTED, size=8)
        add_text(slide, line, 9.13, y + 0.66, 3.1, 0.25, 10, WHITE)

    add_rect(slide, 8.03, 6.55, 4.68, 0.3, PANEL_2, radius=True)
    add_text(
        slide,
        "La Chine influence l’équilibre — sans parité militaire.",
        8.15,
        6.61,
        4.42,
        0.17,
        8,
        MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Sources : OTAN, 2026 ; DoD, 2024 ; CSIS.", 3)
    add_notes(
        slide,
        """
        Trois acteurs structurent le rapport de force.
        La Russie possède la plus longue façade arctique et protège autour de Kola
        ses sous-marins nucléaires : c’est la logique du bastion.
        L’OTAN bénéficie désormais d’une continuité géographique :
        sept des huit États arctiques sont alliés depuis les adhésions finlandaise et suédoise.
        La Chine cherche surtout l’accès aux routes, à la science et aux ressources.
        Son influence augmente, mais elle n’a ni côte arctique ni présence militaire comparable.
        """,
    )

    # 4 — Hotspots
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_kicker(slide, 4, "OÙ EST LE RISQUE ?")
    add_text(slide, "Trois points chauds", 0.62, 1.16, 6.2, 0.56, 29, WHITE, True)
    add_text(
        slide,
        "Un incident local peut produire une réaction régionale.",
        0.65,
        1.75,
        6.4,
        0.32,
        13,
        MUTED,
    )

    hotspots = [
        ("01", "KOLA · GIUK", "Sous-marins\n& renforts", RED),
        ("02", "GROENLAND", "Alerte\n& espace", ICE),
        ("03", "BÉRING · NSR", "Routes\n& zone grise", ORANGE),
    ]
    for i, (num, title, desc, color) in enumerate(hotspots):
        x = 0.63 + i * 4.13
        add_rect(slide, x, 2.35, 3.78, 2.8, PANEL, radius=True)
        add_pill(slide, num, x + 0.25, 2.62, 0.58, fill=color, color=NAVY)
        if i == 0:
            add_circle(slide, x + 1.25, 3.18, 0.96, NAVY_2, line=color, line_width=2)
            add_circle(slide, x + 1.52, 3.45, 0.42, color)
            add_line(slide, x + 2.18, 3.66, x + 3.13, 3.66, color, 2.2)
        elif i == 1:
            add_icon_radar(slide, x + 1.2, 3.14, color)
        else:
            add_icon_route(slide, x + 1.25, 3.57, color)
        add_text(slide, title, x + 0.25, 4.18, 3.28, 0.28, 12, color, True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.25, 4.53, 3.28, 0.46, 11, WHITE, True, align=PP_ALIGN.CENTER)

    add_rect(slide, 0.63, 5.55, 12.03, 0.97, NAVY_2, radius=True, line=PANEL_2)
    chain = [("INCIDENT", RED), ("PERCEPTION", ORANGE), ("RIPOSTE", ICE)]
    for i, (label, color) in enumerate(chain):
        x = 1.08 + i * 3.96
        add_text(slide, label, x, 5.89, 2.25, 0.24, 12, color, True, align=PP_ALIGN.CENTER)
        if i < 2:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(x + 2.56),
                Inches(5.82),
                Inches(0.48),
                Inches(0.4),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(MUTED_2)
            shape.line.fill.background()
    add_text(
        slide,
        "Le risque principal : l’escalade involontaire.",
        0.65,
        6.72,
        12.0,
        0.27,
        12,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Synthèse : OTAN, Conseil de l’Arctique, DoD.", 4)
    add_notes(
        slide,
        """
        Les risques se concentrent dans trois zones.
        Autour de Kola et du passage GIUK, l’enjeu est nucléaire et maritime.
        Au Groenland, ce sont l’alerte avancée, les satellites et l’accès aux infrastructures.
        Dans le détroit de Béring et sur la route maritime du Nord,
        le danger vient des patrouilles, du brouillage et des actions de zone grise.
        Le scénario le plus probable n’est pas une guerre planifiée,
        mais un incident mal interprété, suivi d’une réaction trop rapide.
        """,
    )

    # 5 — Questions
    slide = prs.slides.add_slide(blank)
    add_full_image(slide, QUESTIONS_DARK)
    add_rect(slide, 0, 0, 5.05, 7.5, NAVY, radius=False)
    add_pill(slide, "05 · DISCUSSION", 0.65, 0.52, 1.55)
    add_text(slide, "QUESTIONS ?", 0.65, 1.35, 4.1, 0.7, 37, ICE, True)
    add_text(slide, "À vous de débattre.", 0.67, 2.13, 3.7, 0.38, 15, MUTED)
    questions = [
        ("01", "La dissuasion rend-elle\nl’Arctique plus sûr ?"),
        ("02", "L’OTAN protège-t-elle\nou accélère-t-elle l’escalade ?"),
        ("03", "Peut-on coopérer sur le climat\ntout en rivalisant militairement ?"),
    ]
    for i, (num, text) in enumerate(questions):
        y = 2.9 + i * 1.18
        add_circle(slide, 0.68, y, 0.45, ICE if i == 0 else ORANGE if i == 1 else GREEN)
        add_text(
            slide,
            num,
            0.68,
            y + 0.12,
            0.45,
            0.2,
            9,
            NAVY,
            True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, text, 1.38, y - 0.02, 3.15, 0.67, 14, WHITE, True)
    add_text(slide, "MERCI", 0.68, 6.73, 1.3, 0.24, 10, MUTED_2, True)
    add_photo_credit(slide, "Illustration éditoriale générée · sans valeur documentaire")
    add_notes(
        slide,
        """
        Pour terminer, je vous propose trois questions.
        Premièrement, la dissuasion rend-elle réellement l’Arctique plus sûr ?
        Deuxièmement, le renforcement de l’OTAN protège-t-il la région
        ou accélère-t-il le dilemme de sécurité ?
        Enfin, peut-on encore coopérer sur le climat, la recherche et le secours
        tout en étant en rivalité militaire ?
        Merci pour votre attention.
        """,
    )

    props = prs.core_properties
    props.title = "La militarisation de l’Arctique — version courte"
    props.subject = "Exposé visuel en 5 diapositives"
    props.author = "Cursor"
    props.comments = "Version sombre, visuelle et concise avec notes du présentateur."
    prs.save(OUT)
    print(f"Created: {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
