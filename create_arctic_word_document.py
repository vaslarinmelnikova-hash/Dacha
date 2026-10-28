from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Texte_Expose_Militarisation_Arctique.docx"

NAVY = "123247"
CYAN = "17A9D4"
GREY = "5F6F7A"


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text, size=10.5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    return paragraph


def add_label(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(CYAN)
    return paragraph


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    for name, size, color in [
        ("Title", 27, NAVY),
        ("Subtitle", 12, GREY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 13, CYAN),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LA MILITARISATION DE L’ARCTIQUE · TEXTE COURT")
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(GREY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Document d’accompagnement · octobre 2026")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("La militarisation\nde l’Arctique")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Texte sec de l’exposé · résumé russe · trois questions")

    add_heading(doc, "Очень коротко по-русски", 1)
    add_bullet(
        doc,
        "Милитаризация — это рост числа баз, радаров, учений и военных патрулей.",
    )
    add_bullet(
        doc,
        "Россия защищает ядерные силы на Кольском полуострове; НАТО усиливает общий северный фланг.",
    )
    add_bullet(
        doc,
        "Китай расширяет влияние, но пока не имеет сопоставимого военного присутствия.",
    )
    add_bullet(
        doc,
        "Главный риск — случайный инцидент, неверная интерпретация и быстрая эскалация.",
    )

    doc.add_page_break()

    add_heading(doc, "Texte de l’exposé", 1)

    add_heading(doc, "Diapositive 1 — Introduction", 2)
    add_body(
        doc,
        "Bonjour. Mon exposé porte sur la militarisation de l’Arctique. "
        "La question est de savoir si la fonte des glaces transforme cette région en nouvelle frontière stratégique. "
        "La militarisation désigne l’augmentation des bases, des radars, des exercices et des patrouilles. "
        "Elle ne signifie pas qu’une guerre est certaine, mais que l’Arctique prend une importance militaire croissante.",
    )

    add_heading(doc, "Diapositive 2 — Pourquoi maintenant ?", 2)
    add_body(
        doc,
        "Le changement climatique rend certaines zones plus accessibles. Entre 1979 et 2025, "
        "le minimum de banquise de septembre a diminué de 12,1 % par décennie. "
        "Des routes peuvent donc rester ouvertes plus longtemps, mais les conditions restent difficiles et coûteuses. "
        "L’Arctique est aussi important pour l’alerte radar, l’aviation et les sous-marins nucléaires. "
        "Le climat crée une possibilité d’accès ; la rivalité entre États transforme cette possibilité en enjeu militaire.",
    )

    add_heading(doc, "Diapositive 3 — Les acteurs", 2)
    add_body(
        doc,
        "La Russie possède la plus longue façade arctique. Autour de la péninsule de Kola, "
        "elle protège les sous-marins nucléaires de la flotte du Nord avec des radars, des avions et des missiles. "
        "L’OTAN bénéficie depuis les adhésions de la Finlande et de la Suède d’une continuité géographique : "
        "sept des huit États arctiques sont membres de l’Alliance. "
        "La Chine investit dans la recherche, les routes et l’énergie, mais sa présence militaire reste limitée.",
    )

    add_heading(doc, "Diapositive 4 — Les zones de risque", 2)
    add_body(
        doc,
        "Trois zones sont particulièrement sensibles. Près de Kola et du passage GIUK, "
        "les sous-marins russes rencontrent la surveillance et les routes de renfort de l’OTAN. "
        "Le Groenland est important pour les radars, les satellites et la défense de l’Amérique du Nord. "
        "Le détroit de Béring et la route maritime du Nord concentrent les enjeux de navigation et de contrôle. "
        "Le danger principal est un incident mal interprété, suivi d’une réaction trop rapide.",
    )

    add_heading(doc, "Diapositive 5 — Conclusion", 2)
    add_body(
        doc,
        "L’Arctique est donc un espace de dissuasion, de surveillance et de compétition. "
        "La présence militaire peut empêcher une attaque directe, mais elle augmente aussi le risque d’erreur. "
        "L’objectif principal doit être d’éviter qu’un incident limité devienne une crise majeure.",
    )

    doc.add_page_break()

    add_heading(doc, "Trois questions et réponses courtes", 1)

    add_heading(doc, "1. La dissuasion rend-elle l’Arctique plus sûr ?", 2)
    add_label(doc, "RÉPONSE")
    add_body(
        doc,
        "Elle réduit le risque d’une attaque volontaire, mais elle augmente les conséquences possibles d’une erreur.",
    )
    add_label(doc, "ПО-РУССКИ")
    add_body(doc, "Сдерживание мешает войне, но повышает цену ошибки.", size=9.5)

    add_heading(doc, "2. L’OTAN protège-t-elle ou accélère-t-elle l’escalade ?", 2)
    add_label(doc, "RÉPONSE")
    add_body(
        doc,
        "Elle protège ses membres, mais la Russie peut interpréter son renforcement comme une menace.",
    )
    add_label(doc, "ПО-РУССКИ")
    add_body(doc, "НАТО защищает союзников, но тревожит Россию.", size=9.5)

    add_heading(
        doc,
        "3. Peut-on coopérer sur le climat tout en rivalisant militairement ?",
        2,
    )
    add_label(doc, "RÉPONSE")
    add_body(
        doc,
        "Oui, dans des domaines limités : science, météo, secours, pollution et pêche.",
    )
    add_label(doc, "ПО-РУССКИ")
    add_body(doc, "Ограниченное практическое сотрудничество возможно.", size=9.5)

    properties = doc.core_properties
    properties.title = "Texte court — La militarisation de l’Arctique"
    properties.subject = "Texte oral, résumé russe et trois réponses courtes"
    properties.author = "Cursor"

    doc.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    build_document()
